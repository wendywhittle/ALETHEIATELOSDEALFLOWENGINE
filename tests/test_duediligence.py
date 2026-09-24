"""Due-diligence tracker tests: checklist seeding, backfill idempotency,
status/notes workflow, document upload/download round-trip, progress counts,
pipeline indicator, detail page, and the Excel Due Diligence sheet."""

import io
import os
import tempfile
import unittest

from openpyxl import load_workbook

import excel_gen
import records
from app import app


def sample(**overrides):
    raw = {"name": "DD Deal", "asset_type": "Multi-family",
           "location": "Vancouver, WA", "purchase_price": 1_000_000,
           "scheduled_rent_annual": 100_000, "vacancy_loss_annual": 5_000,
           "other_income_annual": 2_000, "operating_expenses": 30_000}
    raw.update(overrides)
    return raw


class DDCase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        os.environ["DEALFLOW_DB"] = os.path.join(self.tmp.name, "test.db")
        os.environ["DEALFLOW_ARTIFACTS"] = os.path.join(self.tmp.name, "art")
        app.config["TESTING"] = True
        self.client = app.test_client()
        self.deal_id = records.submit_deal(sample())

    def tearDown(self):
        del os.environ["DEALFLOW_DB"]
        del os.environ["DEALFLOW_ARTIFACTS"]
        self.tmp.cleanup()

    def test_checklist_seeded_on_deal_creation(self):
        items = records.list_dd_items(self.deal_id)
        self.assertEqual(len(items), len(records.DD_CHECKLIST))
        self.assertEqual(len(items), 19)
        categories = []
        for i in items:
            if i["category"] not in categories:
                categories.append(i["category"])
        self.assertEqual(categories, records.DD_CATEGORIES)
        self.assertIn("Legal / Title", categories)
        self.assertIn("Financing", categories)
        keys = {i["item_key"] for i in items}
        self.assertIn("hoa_docs", keys)
        self.assertIn("phase_i", keys)
        for i in items:
            self.assertEqual(i["status"], "not_started")
            self.assertEqual(i["notes"], "")

    def test_backfill_seeds_without_duplicating(self):
        # Simulate a legacy deal: wipe its checklist, then backfill twice.
        conn = records._connect()
        conn.execute("DELETE FROM dd_items WHERE deal_id = ?",
                     (self.deal_id,))
        conn.commit()
        conn.close()
        self.assertEqual(records.list_dd_items(self.deal_id), [])
        n = records.backfill_dd_checklist()
        self.assertGreaterEqual(n, 1)
        self.assertEqual(len(records.list_dd_items(self.deal_id)), 19)
        records.backfill_dd_checklist()
        self.assertEqual(len(records.list_dd_items(self.deal_id)), 19)

    def test_backfill_does_not_reset_existing_work(self):
        records.set_dd_status(self.deal_id, "inspection", "cleared")
        records.set_dd_notes(self.deal_id, "inspection", "passed 2026-09-01")
        records.backfill_dd_checklist()
        items = {i["item_key"]: i for i in records.list_dd_items(self.deal_id)}
        self.assertEqual(items["inspection"]["status"], "cleared")
        self.assertEqual(items["inspection"]["notes"], "passed 2026-09-01")

    def test_status_change_persists(self):
        records.set_dd_status(self.deal_id, "appraisal", "requested")
        items = {i["item_key"]: i for i in records.list_dd_items(self.deal_id)}
        self.assertEqual(items["appraisal"]["status"], "requested")
        # any transition allowed
        records.set_dd_status(self.deal_id, "appraisal", "cleared")
        items = {i["item_key"]: i for i in records.list_dd_items(self.deal_id)}
        self.assertEqual(items["appraisal"]["status"], "cleared")

    def test_invalid_status_rejected(self):
        with self.assertRaises(ValueError):
            records.set_dd_status(self.deal_id, "appraisal", "bogus")
        with self.assertRaises(ValueError):
            records.set_dd_status(self.deal_id, "no_such_item", "cleared")

    def test_notes_saved(self):
        records.set_dd_notes(self.deal_id, "title_commitment",
                             "exceptions 4 and 7 need review")
        items = {i["item_key"]: i for i in records.list_dd_items(self.deal_id)}
        self.assertEqual(items["title_commitment"]["notes"],
                         "exceptions 4 and 7 need review")

    def test_progress_count_correct(self):
        cleared, total = records.dd_progress(self.deal_id)
        self.assertEqual((cleared, total), (0, 19))
        records.set_dd_status(self.deal_id, "inspection", "cleared")
        records.set_dd_status(self.deal_id, "appraisal", "cleared")
        records.set_dd_status(self.deal_id, "hoa_docs", "waived")
        cleared, total = records.dd_progress(self.deal_id)
        self.assertEqual((cleared, total), (2, 19))

    def test_upload_download_round_trip(self):
        payload = b"%PDF-1.4 fake inspection report bytes"
        r = self.client.post(
            "/deals/%s/dd/upload" % self.deal_id,
            data={"item_key": "inspection", "note": "full report",
                  "document": (io.BytesIO(payload), "inspection.pdf")},
            content_type="multipart/form-data")
        self.assertEqual(r.status_code, 302)
        docs = records.list_dd_documents(self.deal_id, "inspection")
        self.assertEqual(len(docs), 1)
        self.assertEqual(docs[0]["filename"], "inspection.pdf")
        self.assertEqual(docs[0]["note"], "full report")
        r2 = self.client.get(
            "/deals/%s/dd/docs/%d" % (self.deal_id, docs[0]["id"]))
        self.assertEqual(r2.status_code, 200)
        self.assertEqual(r2.data, payload)

    def test_upload_missing_file_rejected(self):
        r = self.client.post("/deals/%s/dd/upload" % self.deal_id,
                             data={"item_key": "inspection"},
                             content_type="multipart/form-data")
        self.assertEqual(r.status_code, 400)

    def test_download_wrong_deal_rejected(self):
        other = records.submit_deal(sample(name="Other Deal"))
        payload = b"secret bytes"
        self.client.post(
            "/deals/%s/dd/upload" % self.deal_id,
            data={"item_key": "inspection",
                  "document": (io.BytesIO(payload), "x.pdf")},
            content_type="multipart/form-data")
        doc = records.list_dd_documents(self.deal_id)[0]
        r = self.client.get("/deals/%s/dd/docs/%d" % (other, doc["id"]))
        self.assertEqual(r.status_code, 404)

    def test_detail_page_shows_dd(self):
        records.set_dd_status(self.deal_id, "inspection", "cleared")
        records.set_dd_status(self.deal_id, "appraisal", "cleared")
        r = self.client.get("/deals/%s" % self.deal_id)
        self.assertEqual(r.status_code, 200)
        self.assertIn(b"Due diligence", r.data)
        self.assertIn(b"Due diligence: 2 of 19 cleared", r.data)
        self.assertIn(b"Property inspection", r.data)

    def test_dd_status_change_via_web(self):
        r = self.client.post("/deals/%s/dd/status" % self.deal_id,
                             data={"item_key": "survey", "status": "received"})
        self.assertEqual(r.status_code, 302)
        items = {i["item_key"]: i for i in records.list_dd_items(self.deal_id)}
        self.assertEqual(items["survey"]["status"], "received")

    def test_pipeline_shows_dd_progress(self):
        records.set_dd_status(self.deal_id, "inspection", "cleared")
        r = self.client.get("/pipeline")
        self.assertEqual(r.status_code, 200)
        self.assertIn(b"1/19", r.data)

    def test_excel_dd_sheet(self):
        records.set_dd_status(self.deal_id, "inspection", "cleared")
        records.set_dd_notes(self.deal_id, "inspection", "passed")
        deal = records.get_deal(self.deal_id)
        deal["dd_items"] = records.list_dd_items(self.deal_id)
        path = excel_gen.generate_excel(deal)
        wb = load_workbook(path, data_only=True)
        try:
            self.assertIn("Due Diligence", wb.sheetnames)
            ws = wb["Due Diligence"]
            rows = list(ws.iter_rows(min_row=3, max_col=4, values_only=True))
            self.assertEqual(len(rows), 19)
            by_label = {r[1]: r for r in rows}
            self.assertEqual(by_label["Property inspection"][2], "cleared")
            self.assertEqual(by_label["Property inspection"][3], "passed")
            self.assertEqual(by_label["Appraisal"][0], "Financing")
        finally:
            wb.close()


if __name__ == "__main__":
    unittest.main()
