"""Excel artifact tests: workbook generated, values match the canonical
record, calculations match the application's underwriting."""

import os
import tempfile
import unittest

from openpyxl import load_workbook

import excel_gen
import records


def cell_by_label(ws, label):
    for row in ws.iter_rows(min_row=1, max_col=1):
        if row[0].value == label:
            return ws.cell(row=row[0].row, column=2).value
    raise AssertionError("label %r not found in %s" % (label, ws.title))


class ExcelCase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        os.environ["DEALFLOW_DB"] = os.path.join(self.tmp.name, "test.db")
        os.environ["DEALFLOW_ARTIFACTS"] = os.path.join(self.tmp.name, "art")
        deal_id = records.submit_deal({
            "name": "Excel Deal", "asset_type": "Multi-family",
            "location": "Vancouver, WA", "purchase_price": 10_000_000,
            "noi": 650_000, "ltv": 70, "interest_rate": 6.5,
            "amortization_years": 25, "hold_period": 5, "exit_cap_rate": 6.0,
        })
        self.deal = records.get_deal(deal_id)
        self.path = excel_gen.generate_excel(self.deal)
        self.wb = load_workbook(self.path, data_only=True)

    def tearDown(self):
        self.wb.close()
        del os.environ["DEALFLOW_DB"]
        del os.environ["DEALFLOW_ARTIFACTS"]
        self.tmp.cleanup()

    def test_workbook_generated_with_required_sheets(self):
        self.assertTrue(os.path.exists(self.path))
        self.assertEqual(
            set(self.wb.sheetnames),
            {"Deal Summary", "Operating Inputs", "Financing", "Returns",
             "Assumptions - Missing Data", "Due Diligence"})

    def test_values_match_canonical_record(self):
        ws = self.wb["Deal Summary"]
        self.assertEqual(cell_by_label(ws, "Deal ID"), self.deal["deal_id"])
        self.assertEqual(cell_by_label(ws, "Deal name"), "Excel Deal")
        self.assertEqual(cell_by_label(ws, "Purchase price"), 10_000_000)
        self.assertEqual(cell_by_label(ws, "NOI (annual)"), 650_000)

    def test_calculations_match_application(self):
        ws = self.wb["Returns"]
        app_uw = self.deal["underwriting"]
        self.assertAlmostEqual(cell_by_label(ws, "Cap rate"),
                               app_uw["cap_rate"]["value"])
        self.assertAlmostEqual(cell_by_label(ws, "DSCR"),
                               app_uw["dscr"]["value"])
        self.assertAlmostEqual(cell_by_label(ws, "IRR (annual)"),
                               app_uw["irr"]["value"])
        ws = self.wb["Financing"]
        self.assertAlmostEqual(cell_by_label(ws, "Loan amount"),
                               app_uw["loan_amount"]["value"])

    def test_missing_data_marked(self):
        ws = self.wb["Assumptions - Missing Data"]
        states = {}
        for row in ws.iter_rows(min_row=3, max_col=2, values_only=True):
            states[row[0]] = row[1]
        self.assertEqual(states["EGI (annual)"], "unavailable")
        self.assertEqual(states["Purchase price"], "supplied")
        self.assertEqual(states["Cap rate"], "calculated")

    def test_rent_detail_and_derived_egi_in_operating_inputs(self):
        deal_id = records.submit_deal({
            "name": "Rent Detail Deal", "asset_type": "Multi-family",
            "location": "Vancouver, WA", "purchase_price": 1_000_000,
            "scheduled_rent_annual": 100_000, "vacancy_loss_annual": 5_000,
            "other_income_annual": 2_000, "operating_expenses": 30_000,
        })
        deal = records.get_deal(deal_id)
        path = excel_gen.generate_excel(deal)
        wb = load_workbook(path, data_only=True)
        try:
            ws = wb["Operating Inputs"]
            self.assertEqual(cell_by_label(ws, "Scheduled rent (annual)"),
                             100_000)
            self.assertEqual(cell_by_label(ws, "Vacancy / collection loss (annual)"),
                             5_000)
            self.assertEqual(cell_by_label(ws, "Other income (annual)"),
                             2_000)
            self.assertEqual(cell_by_label(ws, "EGI (annual) — result"),
                             97_000)
            states = {}
            for row in ws.iter_rows(min_row=3, max_col=3, values_only=True):
                states[row[0]] = row[2]
            self.assertEqual(states["EGI (annual) — result"], "calculated")
        finally:
            wb.close()

    def test_dd_sheet_lists_checklist(self):
        deal = records.get_deal(self.deal["deal_id"])
        deal["dd_items"] = records.list_dd_items(self.deal["deal_id"])
        path = excel_gen.generate_excel(deal)
        wb = load_workbook(path, data_only=True)
        try:
            ws = wb["Due Diligence"]
            rows = list(ws.iter_rows(min_row=3, max_col=4, values_only=True))
            self.assertEqual(len(rows), 19)
            self.assertEqual(rows[0][0], "Legal / Title")
            self.assertEqual(rows[0][1], "Purchase agreement")
            self.assertEqual(rows[0][2], "not_started")
        finally:
            wb.close()


if __name__ == "__main__":
    unittest.main()
